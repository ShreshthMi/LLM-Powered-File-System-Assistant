"""Helpers that build real PDF and DOCX fixtures for the tests."""

from __future__ import annotations

from pathlib import Path

import docx
from docx.oxml import parse_xml
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

_NAMESPACES = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:v="urn:schemas-microsoft-com:vml"'
)


def make_pdf(path: Path, pages: list[list[str]], *, user_password: str | None = None) -> Path:
    """Write a PDF with one list of text lines per page (an empty list makes a blank page).

    With ``user_password`` the file is AES-256 encrypted; ``""`` gives the common
    "permissions only" protection that opens without a password.
    """
    pdf = canvas.Canvas(str(path))
    for lines in pages:
        for index, line in enumerate(lines):
            pdf.drawString(72, 720 - 18 * index, line)
        pdf.showPage()
    pdf.save()
    if user_password is not None:
        writer = PdfWriter(clone_from=str(path))
        writer.encrypt(user_password=user_password, owner_password="owner-secret", algorithm="AES-256")
        writer.write(str(path))
    return path


def make_docx(
    path: Path,
    paragraphs: list[str],
    *,
    header: str | None = None,
    table: list[list[str]] | None = None,
    raw_body_xml: tuple[str, ...] = (),
) -> Path:
    """Write a DOCX with paragraphs, an optional header and table, and optional raw body XML."""
    document = docx.Document()
    if header is not None:
        document.sections[0].header.paragraphs[0].text = header
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for column_index, value in enumerate(row):
                grid.cell(row_index, column_index).text = value
    body = document.element.body
    for xml in raw_body_xml:
        body.insert(len(body) - 1, parse_xml(xml))  # keep w:sectPr last
    document.save(str(path))
    return path


def content_control_xml(text: str) -> str:
    """A body-level content control (w:sdt), as used by Word's resume templates."""
    return (
        f"<w:sdt {_NAMESPACES}><w:sdtPr/><w:sdtContent>"
        f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"
        "</w:sdtContent></w:sdt>"
    )


def text_box_xml(text: str) -> str:
    """A paragraph holding a text box stored twice: DrawingML plus the legacy VML fallback."""
    box = f"<w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent>"
    return (
        f"<w:p {_NAMESPACES}><w:r><mc:AlternateContent>"
        f'<mc:Choice Requires="wps"><w:drawing><wps:txbx>{box}</wps:txbx></w:drawing></mc:Choice>'
        f"<mc:Fallback><w:pict><v:textbox>{box}</v:textbox></w:pict></mc:Fallback>"
        "</mc:AlternateContent></w:r></w:p>"
    )
