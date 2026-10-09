"""Generate the sample resumes in ``resumes/`` (3 PDF, 3 DOCX, 2 TXT).

Every person and company is fictional; emails use example.com and phone numbers use the
555-01xx range reserved for fiction. The content is chosen to exercise the assistant:

* about half the resumes mention Python, spelled differently each time ("Django/Python",
  "python 3", "PYTHON", "Python/NumPy");
* two people are called John, for "ambiguous file" questions;
* one resume says "JavaScript" but never "Java", and another mentions Java only in
  passing, for whole-word search;
* one mentions "Monty Python" as a hobby, which is not Python experience;
* one DOCX keeps its contact details in the page header and its skills in a table;
* one PDF runs to two pages.

Run from the project folder:  python scripts/generate_resumes.py
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

import docx
from reportlab import rl_config
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate

OUTPUT_DIR = Path(__file__).resolve().parent.parent / "resumes"
FIXED_DATE = datetime(2026, 1, 15, 9, 0, 0)  # stable document metadata between runs

rl_config.invariant = 1  # reproducible PDFs: no timestamps or random IDs


@dataclass(frozen=True)
class Job:
    role: str
    company: str
    dates: str
    highlights: tuple[str, ...]


@dataclass(frozen=True)
class Resume:
    filename: str
    name: str
    title: str
    email: str
    phone: str
    location: str
    summary: str
    skills: tuple[str, ...]
    jobs: tuple[Job, ...]
    education: tuple[str, ...]
    extras: tuple[tuple[str, tuple[str, ...]], ...] = ()  # (section heading, lines)
    skills_table: tuple[tuple[str, str], ...] = ()  # DOCX only: (area, tools) rows
    contact_in_header: bool = False  # DOCX only
    extras_on_new_page: bool = False  # PDF only

    @property
    def contact(self) -> str:
        """Format the fictional email, phone number and location on one line."""
        return f"{self.email} | {self.phone} | {self.location}"


RESUMES: tuple[Resume, ...] = (
    Resume(
        filename="resume_john_doe.pdf",
        name="John Doe",
        title="Senior Backend Engineer",
        email="john.doe@example.com",
        phone="(415) 555-0101",
        location="Austin, TX",
        summary=(
            "Backend engineer with 8 years of experience building high-traffic web services with Django/Python "
            "and FastAPI. Enjoys designing clean APIs, tuning PostgreSQL queries and mentoring junior developers."
        ),
        skills=("Python", "Django", "FastAPI", "Celery", "PostgreSQL", "Redis", "Docker", "AWS (ECS, RDS, S3)",
                "REST and GraphQL APIs", "pytest"),
        jobs=(
            Job("Senior Backend Engineer", "Brightline Logistics", "2021 - Present", (
                "Led the rewrite of the shipment-tracking API in Django/Python, cutting p95 latency from 900 ms to 180 ms.",
                "Designed an event pipeline with Celery and Redis that processes 2 million status updates a day.",
                "Mentored four engineers and introduced contract testing with pytest.",
            )),
            Job("Software Engineer", "Pinecrest Software", "2018 - 2021", (
                "Built billing microservices in Python 3 and FastAPI serving 40,000 customers.",
                "Moved slow reports to PostgreSQL materialized views, saving 30 hours of manual work a month.",
            )),
            Job("Junior Developer", "Orbitflow", "2016 - 2018", (
                "Maintained internal tools and wrote data-import scripts for the operations team.",
            )),
        ),
        education=("B.S. Computer Science, Westbrook University, 2016", "AWS Certified Developer - Associate"),
    ),
    Resume(
        filename="resume_priya_sharma.pdf",
        name="Priya Sharma",
        title="Data Scientist",
        email="priya.sharma@example.com",
        phone="(415) 555-0102",
        location="Seattle, WA",
        summary=(
            "Data scientist with 6 years of experience turning messy data into decisions. Builds forecasting and "
            "churn models in python 3 (pandas, scikit-learn, statsmodels) and explains results to non-technical "
            "stakeholders."
        ),
        skills=("python 3", "pandas", "scikit-learn", "statsmodels", "SQL (Snowflake, BigQuery)", "Tableau",
                "A/B testing", "causal inference", "Airflow"),
        jobs=(
            Job("Senior Data Scientist", "Cobalt Health Labs", "2022 - Present", (
                "Built a patient no-show prediction model (AUC 0.87) that reduced missed appointments by 18%.",
                "Designed A/B tests for the reminder service and automated weekly KPI reports with Airflow.",
            )),
            Job("Data Analyst, then Data Scientist", "Harborview Retail", "2019 - 2022", (
                "Forecast demand for 3,000 products with gradient-boosted models, reducing stock-outs by 12%.",
                "Created Tableau dashboards used every week by more than 200 store managers.",
            )),
        ),
        education=("M.S. Statistics, Lakeshore Institute of Technology, 2019",
                   "B.Sc. Mathematics, Riverside State University, 2017"),
        extras=(
            ("Publications and Talks", (
                "Forecasting retail demand with hierarchical models - Regional Data Science Meetup, 2023.",
                "Practical causal inference for product teams - internal workshop series, 2024.",
            )),
            ("Volunteering", ("Teaches a free introductory statistics workshop at a community centre.",)),
        ),
        extras_on_new_page=True,
    ),
    Resume(
        filename="resume_marcus_chen.pdf",
        name="Marcus Chen",
        title="Java Developer",
        email="marcus.chen@example.com",
        phone="(415) 555-0103",
        location="Chicago, IL",
        summary=(
            "Java developer with 7 years of experience building reliable microservices for payments and banking. "
            "Works across the JVM stack with a focus on clean architecture and observability."
        ),
        skills=("Java 17", "Spring Boot", "Spring Cloud", "Hibernate/JPA", "Apache Kafka", "PostgreSQL", "Oracle",
                "Docker", "Kubernetes", "JUnit 5", "Maven", "Gradle"),
        jobs=(
            Job("Senior Java Developer", "Lumen Ridge Bank", "2020 - Present", (
                "Built a Kafka-based payment-event platform in Java 17 and Spring Boot handling 5,000 transactions a second.",
                "Reduced settlement errors by 40% with idempotent consumers and the outbox pattern.",
            )),
            Job("Java Developer", "Atlas Peak Insurance", "2017 - 2020", (
                "Developed policy-management REST services with Spring Boot and Hibernate.",
                "Raised unit-test coverage from 45% to 85% with JUnit 5 and Testcontainers.",
            )),
        ),
        education=("B.Eng. Software Engineering, Riverside State University, 2017",),
    ),
    Resume(
        filename="resume_john_smith.docx",
        name="John Smith",
        title="Frontend Engineer",
        email="john.smith@example.com",
        phone="(415) 555-0104",
        location="Denver, CO",
        summary=(
            "Frontend engineer with 5 years of experience building accessible, fast web apps with JavaScript, "
            "TypeScript and React. Cares about design systems, performance budgets and developer experience."
        ),
        skills=("JavaScript (ES2023)", "TypeScript", "React", "Next.js", "Redux Toolkit", "Node.js", "HTML5",
                "CSS and Sass", "Jest", "Playwright", "Figma", "WCAG 2.2"),
        jobs=(
            Job("Frontend Engineer", "Quillfeather Media", "2022 - Present", (
                "Rebuilt the subscriber dashboard in React and Next.js, improving Largest Contentful Paint from 4.1 s to 1.6 s.",
                "Created a shared component library now used by six product teams.",
            )),
            Job("Web Developer", "Mosaic Learning", "2020 - 2022", (
                "Built interactive course pages in JavaScript and TypeScript.",
                "Introduced Playwright end-to-end tests that caught more than 30 regressions before release.",
            )),
        ),
        education=("B.A. Interactive Media, Westbrook University, 2020",),
    ),
    Resume(
        filename="resume_elena_rodriguez.docx",
        name="Elena Rodriguez",
        title="DevOps / Site Reliability Engineer",
        email="elena.rodriguez@example.com",
        phone="(415) 555-0105",
        location="Remote (Lisbon)",
        summary=(
            "Site reliability engineer with 9 years of experience keeping production boring. Automates everything: "
            "infrastructure as code with Terraform, release tooling in Go, and operational scripts in PYTHON and Bash."
        ),
        skills=(),
        skills_table=(
            ("Cloud", "AWS, GCP"),
            ("Infrastructure as code", "Terraform, Ansible"),
            ("Containers", "Docker, Kubernetes, Helm"),
            ("Languages", "Go, PYTHON, Bash"),
            ("Observability", "Prometheus, Grafana, OpenTelemetry"),
        ),
        jobs=(
            Job("Staff Site Reliability Engineer", "Tidewater Robotics", "2021 - Present", (
                "Cut monthly cloud spend by 27% by right-sizing Kubernetes clusters.",
                "Wrote PYTHON automation that rotates credentials across 300 services with zero downtime.",
                "Leads incident response and blameless postmortems.",
            )),
            Job("DevOps Engineer", "Saffron Grid Energy", "2016 - 2021", (
                "Migrated 120 virtual machines to Kubernetes with Terraform and Helm.",
                "Built CI/CD pipelines that reduced deploy time from 45 to 8 minutes.",
            )),
        ),
        education=("B.Sc. Computer Engineering, Lakeshore Institute of Technology, 2015",),
        contact_in_header=True,
    ),
    Resume(
        filename="resume_aisha_bello.docx",
        name="Aisha Bello",
        title="Senior Product Manager",
        email="aisha.bello@example.com",
        phone="(415) 555-0106",
        location="New York, NY",
        summary=(
            "Product manager with 7 years of experience shipping B2B SaaS products. Turns customer research into "
            "clear roadmaps and works closely with engineering, design and sales."
        ),
        skills=("Product strategy", "Roadmapping", "User research", "A/B testing", "SQL", "Amplitude", "Jira",
                "Figma", "Pricing and packaging", "Stakeholder management"),
        jobs=(
            Job("Senior Product Manager", "Granite Cloud", "2021 - Present", (
                "Launched a self-serve analytics tier that added $4.2M in annual recurring revenue in its first year.",
                "Partnered with the Java backend team to redesign the permissions model, cutting enterprise onboarding from 3 weeks to 4 days.",
            )),
            Job("Product Manager", "Bluebird Payroll", "2018 - 2021", (
                "Owned the payroll-approvals workflow used by 12,000 small businesses.",
                "Ran more than 25 A/B tests on onboarding, lifting activation by 9 points.",
            )),
        ),
        education=("MBA, Riverside State University, 2018", "B.A. Economics, Westbrook University, 2014"),
    ),
    Resume(
        filename="resume_david_kim.txt",
        name="David Kim",
        title="Machine Learning Engineer",
        email="david.kim@example.com",
        phone="(415) 555-0107",
        location="San Francisco, CA",
        summary=(
            "Machine learning engineer with 6 years of experience taking models from notebook to production. Writes "
            "Python/NumPy and PyTorch code that is tested, profiled and monitored."
        ),
        skills=("Python/NumPy", "PyTorch", "Hugging Face Transformers", "ONNX", "MLflow", "Kubeflow", "Docker",
                "FastAPI", "SQL", "Apache Spark"),
        jobs=(
            Job("Machine Learning Engineer", "Orbitflow", "2021 - Present", (
                "Deployed a real-time fraud model (PyTorch exported to ONNX) serving 3,000 predictions a second at 15 ms p99.",
                "Built an MLflow model registry with automated drift monitoring.",
            )),
            Job("Data Engineer", "Pinecrest Software", "2019 - 2021", (
                "Built Spark pipelines for the recommendation system, cutting the nightly batch from 6 hours to 90 minutes.",
            )),
        ),
        education=("M.S. Computer Science (Machine Learning), Lakeshore Institute of Technology, 2019",),
    ),
    Resume(
        filename="resume_sofia_novak.txt",
        name="Sofia Novak",
        title="Mobile Engineer (iOS and Android)",
        email="sofia.novak@example.com",
        phone="(415) 555-0108",
        location="Toronto, ON",
        summary=(
            "Mobile engineer with 6 years of experience shipping native apps in Swift and Kotlin. Focused on smooth "
            "UX, offline-first data sync and crash-free releases."
        ),
        skills=("Swift", "SwiftUI", "Kotlin", "Jetpack Compose", "Coroutines", "Core Data", "Room", "Firebase",
                "GraphQL", "XCTest", "Espresso", "Fastlane"),
        jobs=(
            Job("Senior Mobile Engineer", "Kestrel Mobile", "2021 - Present", (
                "Led the rewrite of the iOS app in SwiftUI; crash-free sessions rose from 97.1% to 99.8%.",
                "Built offline sync for the Android app with Kotlin coroutines and Room.",
            )),
            Job("Mobile Developer", "Mosaic Learning", "2018 - 2021", (
                "Built the flashcard app (more than 1 million downloads) for iOS and Android.",
            )),
        ),
        education=("B.Sc. Computer Science, Riverside State University, 2018",),
        extras=(("Interests", ("Rock climbing, film photography, and quoting Monty Python at inappropriate moments.",)),),
    ),
)


# ------------------------------------------------------------------------- PDF


def _pdf_styles() -> dict[str, ParagraphStyle]:
    """Construct the paragraph styles used by generated PDF resumes."""
    base = getSampleStyleSheet()
    return {
        "name": ParagraphStyle("Name", parent=base["Title"], fontSize=20, spaceAfter=4),
        "title": ParagraphStyle("Role", parent=base["Heading3"], alignment=1, spaceAfter=2),
        "contact": ParagraphStyle("Contact", parent=base["BodyText"], alignment=1, spaceAfter=10),
        "section": ParagraphStyle("Section", parent=base["Heading2"], spaceBefore=10, spaceAfter=4),
        "job": ParagraphStyle("Job", parent=base["Heading4"], spaceBefore=6, spaceAfter=2),
        "body": base["BodyText"],
    }


def _bullets(lines: tuple[str, ...], style: ParagraphStyle) -> ListFlowable:
    """Build an escaped, styled PDF bullet list."""
    items = [ListItem(Paragraph(escape(line), style), leftIndent=12) for line in lines]
    return ListFlowable(items, bulletType="bullet", leftIndent=12)


def write_pdf(resume: Resume, path: Path) -> None:
    """Render ``resume`` as a one- or two-page PDF with reportlab."""
    styles = _pdf_styles()
    story = [
        Paragraph(escape(resume.name), styles["name"]),
        Paragraph(escape(resume.title), styles["title"]),
        Paragraph(escape(resume.contact), styles["contact"]),
        Paragraph("Summary", styles["section"]),
        Paragraph(escape(resume.summary), styles["body"]),
        Paragraph("Skills", styles["section"]),
        Paragraph(escape(", ".join(resume.skills)), styles["body"]),
        Paragraph("Experience", styles["section"]),
    ]
    for job in resume.jobs:
        story += [Paragraph(escape(f"{job.role} - {job.company} ({job.dates})"), styles["job"]),
                  _bullets(job.highlights, styles["body"])]
    story += [Paragraph("Education", styles["section"]), _bullets(resume.education, styles["body"])]
    if resume.extras_on_new_page:
        story.append(PageBreak())
    for heading, lines in resume.extras:
        story += [Paragraph(escape(heading), styles["section"]), _bullets(lines, styles["body"])]
    document = SimpleDocTemplate(str(path), pagesize=LETTER, title=f"{resume.name} - Resume", author=resume.name,
                                 topMargin=0.7 * inch, bottomMargin=0.7 * inch)
    document.build(story)


# ------------------------------------------------------------------------ DOCX


def write_docx(resume: Resume, path: Path) -> None:
    """Render ``resume`` as a Word document with python-docx."""
    document = docx.Document()
    properties = document.core_properties
    properties.author, properties.title = resume.name, f"{resume.name} - Resume"
    properties.created = properties.modified = FIXED_DATE
    if resume.contact_in_header:
        document.sections[0].header.paragraphs[0].text = f"{resume.name} | {resume.contact}"
    document.add_heading(resume.name, level=0)
    document.add_paragraph(resume.title)
    if not resume.contact_in_header:
        document.add_paragraph(resume.contact)
    document.add_heading("Summary", level=1)
    document.add_paragraph(resume.summary)
    document.add_heading("Skills", level=1)
    if resume.skills_table:
        table = document.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Area", "Tools"
        for area, tools in resume.skills_table:
            row = table.add_row().cells
            row[0].text, row[1].text = area, tools
    else:
        document.add_paragraph(", ".join(resume.skills))
    document.add_heading("Experience", level=1)
    for job in resume.jobs:
        document.add_heading(f"{job.role} - {job.company} ({job.dates})", level=2)
        for line in job.highlights:
            document.add_paragraph(line, style="List Bullet")
    document.add_heading("Education", level=1)
    for line in resume.education:
        document.add_paragraph(line, style="List Bullet")
    for heading, lines in resume.extras:
        document.add_heading(heading, level=1)
        for line in lines:
            document.add_paragraph(line, style="List Bullet")
    document.save(str(path))


# ------------------------------------------------------------------------- TXT


def _txt_section(heading: str, lines: list[str]) -> list[str]:
    """Format a plain-text section with an uppercase heading."""
    return ["", heading.upper(), "-" * len(heading), *lines]


def write_txt(resume: Resume, path: Path) -> None:
    """Render ``resume`` as plain UTF-8 text."""
    lines = [resume.name.upper(), resume.title, resume.contact]
    lines += _txt_section("Summary", [resume.summary])
    lines += _txt_section("Skills", [", ".join(resume.skills)])
    experience: list[str] = []
    for job in resume.jobs:
        experience += [f"{job.role} - {job.company} ({job.dates})", *[f"  * {item}" for item in job.highlights], ""]
    lines += _txt_section("Experience", experience[:-1])
    lines += _txt_section("Education", [f"  * {item}" for item in resume.education])
    for heading, extra in resume.extras:
        lines += _txt_section(heading, list(extra))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


WRITERS = {".pdf": write_pdf, ".docx": write_docx, ".txt": write_txt}


def generate_all(output_dir: Path = OUTPUT_DIR) -> list[Path]:
    """Write every sample resume into ``output_dir`` and return the paths."""
    output_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for resume in RESUMES:
        path = output_dir / resume.filename
        WRITERS[path.suffix](resume, path)
        written.append(path)
    return written


def main() -> None:
    """Generate the sample resumes and list them."""
    for path in generate_all():
        print(f"wrote {path.parent.name}/{path.name} ({path.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
